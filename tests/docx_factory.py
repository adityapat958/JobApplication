"""Build fake Word resumes with python-docx (no personal data)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def add_link(paragraph, url, text):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def para(doc, *parts, style=None):
    p = doc.add_paragraph(style=style)
    for part in parts:
        if isinstance(part, tuple) and part[0] == "link":
            add_link(p, part[1], part[2])
        elif isinstance(part, tuple):
            run = p.add_run(part[1])
            run.bold = "b" in part[0]
            run.italic = "i" in part[0]
        else:
            p.add_run(part)
    return p


def simple_resume(path):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Inches(0.7))
    doc.styles["Normal"].font.size = Pt(10.5)
    p = para(doc, ("b", "Sam Sample"))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Finance Student | Oslo, Norway")
    para(doc, ("link", "mailto:sam@example.com", "sam@example.com"), " | ",
         ("link", "https://www.linkedin.com/in/sample", "LinkedIn"))
    doc.add_heading("Education", level=1)
    para(doc, ("b", "Example Business School"), ", Aug 2024 – June 2026")
    para(doc, "MSc in Finance, GPA 4.8/5.0")
    doc.add_heading("Skills", level=1)
    para(doc, ("b", "Core Competencies:"), " Financial Analysis, Valuation, Market Research")
    para(doc, ("b", "Technical Skills:"), " Excel, SQL, Power BI")
    doc.add_heading("Experience", level=1)
    para(doc, ("b", "Fjord Capital"), ", June 2025 – Aug 2025")
    para(doc, ("i", "Summer Analyst"))
    para(doc, "Built a valuation model for 12 Nordic utilities", style="List Bullet")
    para(doc, "Presented findings to the investment committee", style="List Bullet")
    doc.add_heading("Projects", level=1)
    para(doc, ("b", "Housing Dashboard"), " – ", ("link", "https://example.com/dash", "demo"))
    para(doc, "Tracked rental prices across 15 districts", style="List Bullet")
    doc.save(path)
    return path


def _shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def builder_resume(path):
    """Resume-builder style .docx: skills in a table, custom heading style, spacer paragraphs,
    U+2011 hyphens, CO₂ subscript, white-on-dark heading (visible) and invisible tracker text."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import RGBColor
    doc = Document()
    doc.styles.add_style("div_heading", WD_STYLE_TYPE.PARAGRAPH)
    para(doc, ("b", "Sam Sample"))
    para(doc, "Data Analytics | Optimization")
    for _ in range(3):
        para(doc, "\u00a0")                                   # spacer paragraphs
    para(doc, "Oslo | sam\u2011sample\u2011123 on Linkedin.com/in/sam\u2011sample\u2011123")
    h = para(doc, "Summary", style="div_heading")
    h.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)       # white text on a dark band: visible
    _shade(h._p.get_or_add_pPr(), "394D6B")
    para(doc, ("b", "Analyst who models CO"), ("b", "\u2082"), ("b", " transport networks."))
    para(doc, "Skills", style="div_heading")
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("\u2022 Programming: Python, SQL").bold = True
    cell.add_paragraph().add_run("\u2022 Domain: supply\u2011chain analytics").bold = True
    para(doc, "Experience", style="div_heading")
    p = para(doc, ("b", "Analyst"), ", 10/2026 - Present")
    hidden = p.add_run(" secretword")
    hidden.font.hidden = True                                   # w:vanish
    para(doc, "Cut cost by 12% with a MILP model", style="List Bullet")
    t = para(doc, "#TRK#0a1b2c3d-feed")
    t.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)       # white on white: invisible
    doc.save(path)
    return path
