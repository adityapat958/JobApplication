"""Word (.docx) or PDF → LaTeX template import.

The LLM (Claude in Code/Desktop, or an API) only *types*; this module does everything checkable:
extract → render prompt → verify gates (compile, pages, token diff, links, markers) → fix report.

Files live in $JOBAPP_HOME/import/:
  source.docx|source.pdf, meta.json, prompt.md, resume.tex (written by the agent), fix_report.md, build/
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from .. import paths
from ..log import get
from ..tools import which
from . import tex as texmod
from .build import check_markers, tailor
from .keywords import Bank

log = get("import")
MAX_ROUNDS = 3
ALLOWLIST = {"github", "link", "linkedin"}  # labels the template's macros add


class ImportError_(RuntimeError):
    @property
    def user_message(self) -> str:
        return str(self)


def import_dir() -> Path:
    return paths.import_dir()


def candidate_path() -> Path:
    return import_dir() / "resume.tex"


def fix_report_path() -> Path:
    return import_dir() / "fix_report.md"


# ── extraction ────────────────────────────────────────────
def _pandoc(docx: Path, fmt: str) -> str | None:
    exe = which("pandoc")
    if not exe:
        return None
    r = subprocess.run([exe, str(docx), "-t", fmt, "--wrap=none"], capture_output=True, text=True, timeout=120)
    return r.stdout if r.returncode == 0 else None


def _docx_paragraph_md(p, rels) -> str:
    out = []
    for child in p._p.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "r":
            out.append(_run_md(child))
        elif tag == "hyperlink":
            rid = child.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            text = "".join(_run_md(r) for r in child.iter() if r.tag.endswith("}r"))
            url = rels.get(rid)
            out.append(f"[{text}]({url})" if url else text)
    return "".join(out)


def _run_md(r) -> str:
    text = "".join(t.text or "" for t in r.iter() if t.tag.endswith("}t"))
    if not text.strip():
        return text
    rpr = r.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
    bold = rpr is not None and rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}b") is not None
    ital = rpr is not None and rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}i") is not None
    if bold:
        text = f"**{text}**"
    if ital:
        text = f"*{text}*"
    return text


def _docx_fallback(docx: Path) -> tuple[str, str]:
    """(markdown-ish, plain) via python-docx when pandoc is missing."""
    import docx as pydocx
    d = pydocx.Document(str(docx))
    rels = {r.rId: r.target_ref for r in d.part.rels.values() if r.is_external}
    md, plain = [], []
    for p in d.paragraphs:
        text = p.text
        if not text.strip():
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        body = _docx_paragraph_md(p, rels)
        if style.startswith("heading") or style == "title":
            lvl = re.findall(r"\d", style)
            md.append("#" * (int(lvl[0]) if lvl else 1) + " " + body)
        elif "list" in style or p._p.pPr is not None and p._p.pPr.numPr is not None:
            md.append("- " + body)
        else:
            md.append(body)
        plain.append(text)
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            md.append("- " + " | ".join(cells))
            plain.append(" ".join(cells))
    return "\n\n".join(md), "\n".join(plain)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_LIGHT = {None, "", "auto", "FFFFFF"}


def _fill(el) -> str | None:
    shd = el.find(f"{_W}shd") if el is not None else None
    return (shd.get(f"{_W}fill") or "").upper() if shd is not None else None


def _invisible(run) -> bool:
    """Hidden (w:vanish) or white-on-white text — e.g. resume-builder tracking codes."""
    rpr = run.find(f"{_W}rPr")
    if rpr is None:
        return False
    if rpr.find(f"{_W}vanish") is not None or rpr.find(f"{_W}specVanish") is not None:
        return True
    color = rpr.find(f"{_W}color")
    if color is None or (color.get(f"{_W}val") or "").upper() != "FFFFFF":
        return False
    fills = [_fill(rpr)]
    for anc in run.iterancestors():
        tag = anc.tag.rsplit("}", 1)[-1]
        if tag == "p":
            fills.append(_fill(anc.find(f"{_W}pPr")))
        elif tag == "tc":
            fills.append(_fill(anc.find(f"{_W}tcPr")))
    return all((f or "").upper() in _LIGHT for f in fills)  # no dark background behind it


def strip_invisible(docx: Path, out_dir: Path) -> tuple[Path, list[str]]:
    """Copy of the .docx without invisible runs (+ the removed texts). Same file if nothing hidden."""
    from lxml import etree
    with zipfile.ZipFile(docx) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        removed = []
        for r in list(root.iter(f"{_W}r")):
            if _invisible(r):
                text = "".join(t.text or "" for t in r.iter(f"{_W}t"))
                if text.strip():
                    removed.append(text.strip())
                r.getparent().remove(r)
        if not removed:
            return docx, []
        out = out_dir / docx.name
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == "word/document.xml":
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                zo.writestr(item, data)
    return out, removed


def docx_urls(docx: Path) -> set[str]:
    with zipfile.ZipFile(docx) as z:
        names = [n for n in z.namelist() if n.startswith("word/_rels/") and n.endswith(".rels")]
        urls = set()
        for n in names:
            xml = z.read(n).decode("utf-8", "replace")
            for m in re.finditer(r'<Relationship [^>]*>', xml):
                tag = m.group(0)
                if 'TargetMode="External"' in tag and "/hyperlink" in tag:
                    t = re.search(r'Target="([^"]+)"', tag)
                    if t:
                        urls.add(_norm_url(_xml_unescape(t[1])))
    return urls


def _xml_unescape(s: str) -> str:
    return s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')


def _norm_url(u: str) -> str:
    return u.strip().rstrip("/")


def docx_pages(docx: Path) -> int | None:
    try:
        with zipfile.ZipFile(docx) as z:
            xml = z.read("docProps/app.xml").decode("utf-8", "replace")
    except (KeyError, zipfile.BadZipFile):
        return None
    m = re.search(r"<Pages>(\d+)</Pages>", xml)
    return int(m[1]) if m and int(m[1]) > 0 else None


def docx_layout(docx: Path) -> dict:
    import docx as pydocx
    d = pydocx.Document(str(docx))
    out: dict = {}
    if d.sections:
        s = d.sections[0]
        emu = 914400

        def inch(v):
            return round(v / emu, 2) if v is not None else None
        out["page_size_in"] = [inch(s.page_width), inch(s.page_height)]
        out["margins_in"] = {"top": inch(s.top_margin), "bottom": inch(s.bottom_margin),
                             "left": inch(s.left_margin), "right": inch(s.right_margin)}
    try:
        sz = d.styles["Normal"].font.size
        out["base_font_pt"] = sz.pt if sz else None
    except KeyError:
        pass
    sections = []
    for p in d.paragraphs:
        t = p.text.strip()
        style = (p.style.name or "").lower() if p.style is not None else ""
        if not t:
            continue
        if "heading" in style or (t.isupper() and len(t) < 40):
            sections.append(t)
    out["section_order"] = sections
    out["tables"] = len(d.tables)
    out["images"] = len(d.inline_shapes)
    return out


@dataclass
class Extracted:
    markdown: str
    plain: str
    layout: dict
    pages: int
    urls: set[str]
    used_pandoc: bool
    kind: str = "docx"


SUPPORTED = (".docx", ".pdf")


def extract(src: Path) -> Extracted:
    docx = Path(src)
    if not docx.exists():
        raise ImportError_(f"File not found: {docx}")
    suffix = docx.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(docx)
    if suffix != ".docx":
        raise ImportError_("Please give a Word (.docx) or PDF file. Old .doc files: in Word, "
                           "File → Save As → .docx.")
    try:
        zipfile.ZipFile(docx).close()
    except zipfile.BadZipFile:
        raise ImportError_("That file doesn't look like a real .docx (it can't be opened).") from None
    with tempfile.TemporaryDirectory() as tmp:
        clean, hidden = strip_invisible(docx, Path(tmp))
        md = _pandoc(clean, "markdown")
        plain = _pandoc(clean, "plain") if md is not None else None
        used = md is not None and plain is not None
        if not used:
            md, plain = _docx_fallback(clean)
        layout = docx_layout(clean)
    if hidden:
        layout["hidden_text_removed"] = f"{len(hidden)} invisible text run(s) (hidden or white-on-white) — not resume content"
    return Extracted(md, plain, layout, docx_pages(docx) or 1, docx_urls(docx), used)


def _extract_pdf(pdf: Path) -> Extracted:
    from . import pdf_extract as px
    try:
        n = px.pages(pdf)
    except Exception:  # noqa: BLE001 — pypdf raises many types on garbage
        raise ImportError_("That file doesn't look like a real PDF (it can't be opened).") from None
    text, used = px.layout_text(pdf)
    plain = px.strip_bullets(px.plain_text(pdf))
    if len(re.findall(r"\w", plain)) < 20:
        raise ImportError_("This PDF has no selectable text (it's probably a scan or a picture). "
                           "Please use the Word file, or export the PDF again from Word/Google Docs.")
    link_list = px.links(pdf)
    md = px.to_markdown(text, link_list)
    return Extracted(md, plain, px.layout_facts(pdf, md, used), n or 1, {u["url"] for u in link_list}, used, "pdf")


# ── prompt ────────────────────────────────────────────────
def _template_tex() -> str:
    return (paths.examples_profile() / "resume" / "resume.tex").read_text(encoding="utf-8")


SOURCE_NAMES = {"docx": "Word (.docx)", "pdf": "PDF"}


def render_prompt(ex: Extracted, fix_report: str = "", agent_footer: bool = True) -> str:
    raw = (paths.prompts_dir() / "docx_to_latex.md").read_text(encoding="utf-8")
    start = raw.find("## SYSTEM")
    end = raw.find("## Verification gates")
    body = raw[start:end if end > 0 else None].rstrip().rstrip("-").rstrip()
    body = re.sub(r"<!--.*?-->", "", body, flags=re.S)
    import yaml
    layout = yaml.safe_dump(ex.layout, sort_keys=False, allow_unicode=True).strip()
    body = (body.replace("{{TEMPLATE_TEX}}", _template_tex().rstrip())
                .replace("{{SOURCE_KIND}}", SOURCE_NAMES[ex.kind])
                .replace("{{SOURCE_MARKDOWN}}", ex.markdown.strip())
                .replace("{{SOURCE_LAYOUT}}", layout)
                .replace("{{TARGET_PAGES}}", str(ex.pages))
                .replace("{{FIX_REPORT}}", fix_report.strip()))
    if agent_footer:
        body += (f"\n\n---\n## Where to put your answer (agent mode)\n"
                 f"- Write ONLY the LaTeX (no ``` fences) to: `{candidate_path()}`\n"
                 f"- Write the mapping notes YAML to: `{import_dir() / 'mapping_notes.yaml'}`\n"
                 f"- Then run: `jobapp resume import --verify`. If it fails, read `{fix_report_path()}` and fix "
                 f"ONLY the listed problems (max {MAX_ROUNDS} rounds).\n")
    return body + "\n"


def _meta_path() -> Path:
    return import_dir() / "meta.json"


def load_meta() -> dict:
    p = _meta_path()
    if not p.exists():
        raise ImportError_("No import in progress. Start with: jobapp resume import <your resume .docx or .pdf> --prepare")
    return json.loads(p.read_text())


def _save_meta(meta: dict) -> None:
    _meta_path().write_text(json.dumps(meta, indent=2, ensure_ascii=False))


def prepare(docx: str | Path) -> Path:
    docx = Path(docx).expanduser()
    ex = extract(docx)
    d = import_dir()
    d.mkdir(parents=True, exist_ok=True)
    for stale in ("source.docx", "source.pdf", "resume.tex", "fix_report.md", "mapping_notes.yaml"):
        (d / stale).unlink(missing_ok=True)
    shutil.copy2(docx, d / f"source.{ex.kind}")
    (d / "source.md").write_text(ex.markdown, encoding="utf-8")
    (d / "source.txt").write_text(ex.plain, encoding="utf-8")
    _save_meta({"source": docx.name, "target_pages": ex.pages, "urls": sorted(ex.urls), "round": 0,
                "used_pandoc": ex.used_pandoc, "kind": ex.kind})
    p = d / "prompt.md"
    p.write_text(render_prompt(ex), encoding="utf-8")
    return p


# ── verification ──────────────────────────────────────────
_HYPH_BREAK = re.compile(r"(\w+)[-\u2010\u00ad][ \t]*\r?\n[ \t]*(\w+)")


def _prep(text: str) -> str:
    t = unicodedata.normalize("NFKC", text)  # also splits ligatures: ﬁ ﬂ ﬀ ﬃ → fi fl ff ffi
    t = re.sub(r"[\u00ad\u200b\u200c\u200d\ufeff](?![ \t]*\r?\n)", "", t)
    return t.casefold()


def normalize_tokens(text: str) -> Counter:
    """Lowercase word tokens; punctuation, separators (| · •) and bullet glyphs drop out."""
    t = re.sub(r"[^\w]+", " ", _prep(text))
    return Counter(w for w in t.split() if w and w not in ALLOWLIST)


def hyphen_breaks(text: str) -> list[tuple[str, str]]:
    """(left, right) word halves split by a hyphen at a line end ("valu-\nation")."""
    return [(m[1], m[2]) for m in _HYPH_BREAK.finditer(_prep(text))]


def content_diff(source: str, output: str) -> tuple[list[str], list[str]]:
    """Token-multiset diff → (missing from output, extra in output).

    A word hyphenated at a line end ("valu-" / "ation") on one side counts as the whole word on the
    other side — but only for that exact break, so real rewording is still caught. Compounds like
    "real-time" are split on both sides anyway.
    """
    src, out = normalize_tokens(source), normalize_tokens(output)
    missing, extra = src - out, out - src
    for brk, whole_side, parts_side in ((hyphen_breaks(output), missing, extra),
                                        (hyphen_breaks(source), extra, missing)):
        for a, b in brk:
            if whole_side[a + b] > 0 and parts_side[a] > 0 and parts_side[b] > 0 and a != b:
                whole_side[a + b] -= 1
                parts_side[a] -= 1
                parts_side[b] -= 1
    return sorted((+missing).elements()), sorted((+extra).elements())


def pdf_text(pdf: Path) -> str:
    exe = which("pdftotext")
    if exe:
        # -raw: the default mode "dehyphenates" every line ending in "-", gluing real compounds
        # ("fail-" / "closed" -> "failclosed"); content_diff handles line-end hyphens itself.
        r = subprocess.run([exe, "-enc", "UTF-8", "-raw", str(pdf), "-"], capture_output=True, text=True)
        if r.returncode == 0:
            return r.stdout
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(str(pdf)).pages)


def tex_urls(tex: str) -> set[str]:
    nocomment = re.sub(r"(?<!\\)%[^\n]*", "", tex)
    urls = set(re.findall(r"\\href\{([^}]*)\}", nocomment)) | set(re.findall(r"\\url\{([^}]*)\}", nocomment))
    # \projentry{Title}{URL}{stack} — the URL is the 2nd argument
    urls |= {m for m in re.findall(r"\\projentry\{[^}]*\}%?\s*\{([^}]*)\}", nocomment)}
    urls = {u for u in urls if not re.fullmatch(r"#\d", u.strip())}  # macro parameters like \href{#2}
    return {_norm_url(u.replace(r"\#", "#").replace(r"\%", "%").replace(r"\_", "_").replace(r"\&", "&"))
            for u in urls}


def extract_latex(text: str) -> str:
    m = re.search(r"```(?:latex|tex)\s*\n(.*?)```", text, re.S)
    return m[1] if m else text


@dataclass
class Gate:
    name: str
    ok: bool
    detail: str = ""


@dataclass
class Report:
    gates: list[Gate] = field(default_factory=list)
    missing: list[str] = field(default_factory=list)
    extra: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    round: int = 0
    pdf: Path | None = None

    @property
    def ok(self) -> bool:
        return bool(self.gates) and all(g.ok for g in self.gates)

    def summary(self) -> str:
        lines = [("✔ " if g.ok else "✘ ") + g.name + (f": {g.detail}" if g.detail else "") for g in self.gates]
        lines += [f"! {w}" for w in self.warnings]
        if self.ok:
            lines.append("All checks passed. Run 'jobapp resume import --accept' to use it as your resume.")
        elif self.round >= MAX_ROUNDS:
            lines.append(f"Still failing after {self.round} rounds — stop and ask the user / run 'jobapp report-bug'.")
        else:
            lines.append(f"Round {self.round} of {MAX_ROUNDS} failed.")
        return "\n".join(lines)

    def fix_report(self) -> str:
        out = ["### Previous attempt failed verification"]
        for g in self.gates:
            if not g.ok:
                out.append(f"- {g.name}: {g.detail}")
        if self.missing:
            out.append(f"- missing words vs source: {json.dumps(self.missing, ensure_ascii=False)}")
        if self.extra:
            out.append(f"- extra words not in source: {json.dumps(self.extra, ensure_ascii=False)}")
        out.append("Fix ONLY these issues; keep everything else identical to your previous output.")
        return "\n".join(out) + "\n"


def verify_tex(tex: str, source_plain: str, source_urls: set[str], target_pages: int,
               bank: Bank | None = None, workdir: Path | None = None, kind: str = "docx") -> Report:
    rep = Report()
    workdir = Path(workdir or (import_dir() / "build"))
    workdir.mkdir(parents=True, exist_ok=True)
    tex_file = workdir / "resume.tex"
    tex_file.write_text(tex, encoding="utf-8")
    # 1. compiles
    try:
        pdf = texmod.compile_pdf(tex_file, workdir)
        rep.pdf = pdf
        rep.gates.append(Gate("Compiles", True))
        log_text = (workdir / "resume.log").read_text(errors="replace") if (workdir / "resume.log").exists() else ""
        over = [float(x) for x in re.findall(r"Overfull \\hbox \(([\d.]+)pt too wide\)", log_text)]
        if any(o > 5 for o in over):
            rep.warnings.append(f"{sum(o > 5 for o in over)} line(s) stick out into the margin (Overfull hbox > 5pt)")
    except texmod.TexError as e:
        rep.gates.append(Gate("Compiles", False, str(e)))
        pdf = None
    # 2. pages
    if pdf:
        n = texmod.page_count(pdf)
        ok = n <= target_pages
        detail = f"{n} page(s), source has {target_pages}" + ("" if ok else
                 f" — {n - target_pages} page(s) too many: tighten the spacing knobs (never cut content)")
        rep.gates.append(Gate("Page count", ok, detail))
        # 3. content
        missing, extra = content_diff(source_plain, pdf_text(pdf))
        rep.missing, rep.extra = missing, extra
        ok = not missing and not extra
        detail = "identical words" if ok else f"{len(missing)} missing, {len(extra)} extra"
        rep.gates.append(Gate(f"Same words as the {'PDF' if kind == 'pdf' else 'Word'} file", ok, detail))
    # 4. links
    got = tex_urls(tex)
    want = {_norm_url(u) for u in source_urls}
    if got == want:
        rep.gates.append(Gate("Links", True, f"{len(got)} link(s)"))
    else:
        parts = []
        if want - got:
            parts.append("missing " + ", ".join(sorted(want - got)))
        if got - want:
            parts.append("extra " + ", ".join(sorted(got - want)))
        rep.gates.append(Gate("Links", False, "; ".join(parts)))
    # 5. markers + build dry-run
    try:
        check_markers(tex)
        b = bank or Bank.load(paths.examples_profile() / "resume" / "keywords.yaml")
        tailor(tex, b, b.core[:2], b.tech[:2])
        rep.gates.append(Gate("Skills markers", True))
    except ValueError as e:
        rep.gates.append(Gate("Skills markers", False, str(e)))
    return rep


def _user_bank() -> Bank | None:
    try:
        from .. import config
        cfg = config.load()
        return Bank.load(cfg.resolve(cfg.resume.keywords))
    except Exception:  # noqa: BLE001
        return None


def verify() -> Report:
    meta = load_meta()
    d = import_dir()
    cand = candidate_path()
    if not cand.exists():
        raise ImportError_(f"Nothing to check yet: write the LaTeX to {cand} first.")
    tex = extract_latex(cand.read_text(encoding="utf-8"))
    if tex != cand.read_text(encoding="utf-8"):
        cand.write_text(tex, encoding="utf-8")
    rep = verify_tex(tex, (d / "source.txt").read_text(encoding="utf-8"), set(meta["urls"]),
                     int(meta["target_pages"]), _user_bank(), kind=meta.get("kind", "docx"))
    meta["round"] = int(meta.get("round", 0)) + 1
    meta["passed"] = rep.ok
    _save_meta(meta)
    rep.round = meta["round"]
    if rep.ok:
        fix_report_path().unlink(missing_ok=True)
    else:
        fix_report_path().write_text(rep.fix_report(), encoding="utf-8")
    return rep


def accept() -> Path:
    meta = load_meta()
    if not meta.get("passed"):
        raise ImportError_("The imported resume hasn't passed the checks yet. Run 'jobapp resume import --verify'.")
    from .. import config
    from ..db import now
    cfg = config.load()
    dst = cfg.resolve(cfg.resume.tex)
    dst.parent.mkdir(parents=True, exist_ok=True)
    if dst.exists():
        backup = dst.with_name(dst.name + ".bak-" + now()[:19].replace(":", "-"))
        shutil.copy2(dst, backup)
    shutil.copy2(candidate_path(), dst)
    return dst


# ── API mode ──────────────────────────────────────────────
def _call_anthropic(prompt: str, key: str) -> str:
    import os

    import httpx
    model = os.environ.get("JOBAPP_MODEL", "claude-sonnet-4-5")
    system, _, user = prompt.partition("## USER")
    r = httpx.post("https://api.anthropic.com/v1/messages", timeout=300,
                   headers={"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
                   json={"model": model, "max_tokens": 16000, "system": system.replace("## SYSTEM", "").strip(),
                         "messages": [{"role": "user", "content": user.strip()}]})
    r.raise_for_status()
    return "".join(b.get("text", "") for b in r.json().get("content", []))


def _call_gemini(prompt: str, key: str) -> str:
    import os

    import httpx
    model = os.environ.get("JOBAPP_GEMINI_MODEL", "gemini-2.5-pro")
    r = httpx.post(f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent",
                   params={"key": key}, timeout=300,
                   json={"contents": [{"parts": [{"text": prompt}]}]})
    r.raise_for_status()
    parts = r.json()["candidates"][0]["content"]["parts"]
    return "".join(p.get("text", "") for p in parts)


def api_caller():
    from .. import secrets
    k = secrets.get("ANTHROPIC_API_KEY")
    if k:
        return lambda p: _call_anthropic(p, k)
    k = secrets.get("GEMINI_API_KEY")
    if k:
        return lambda p: _call_gemini(p, k)
    return None


def run_api(docx: str | Path, call=None) -> Report:
    call = call or api_caller()
    if call is None:
        raise ImportError_("No ANTHROPIC_API_KEY or GEMINI_API_KEY found. Use agent mode instead: "
                           "jobapp resume import <your resume .docx or .pdf> --prepare")
    prepare(docx)
    ex = extract(Path(docx).expanduser())
    fix = ""
    rep = Report()
    for _ in range(MAX_ROUNDS):
        answer = call(render_prompt(ex, fix, agent_footer=False))
        candidate_path().write_text(extract_latex(answer), encoding="utf-8")
        notes = re.search(r"```yaml\s*\n(.*?)```", answer, re.S)
        if notes:
            (import_dir() / "mapping_notes.yaml").write_text(notes[1], encoding="utf-8")
        rep = verify()
        if rep.ok:
            break
        prev = candidate_path().read_text(encoding="utf-8")
        fix = rep.fix_report() + "\nYour previous output:\n```latex\n" + prev + "\n```\n"
    return rep
